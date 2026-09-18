"""
consolidar_facturas.py
----------------------
Agrega a la planilla anual de facturación los comprobantes listados en los
subdiarios de ventas, completando cada uno con sus ítems desde los .XLS
exportados del sistema.

PROCEDIMIENTO:
  1. Toma la planilla original (correcta hasta julio) como base.
  2. Lee los subdiarios (carpeta 'subdiarios') -> listado de comprobantes con
     sus cabeceras (fecha, tipo, punto de venta, número, CUIT, cliente, total, CAE).
  3. Para cada comprobante, busca su .XLS en la carpeta 'comprobantes' y agrega
     una fila por cada ítem (cabecera + descripción, cantidad, precio, IVA).
  4. Si falta el .XLS de un comprobante, agrega igual la fila con la cabecera
     (detalle vacío) y lo incluye en la lista de faltantes.
  5. Ordena toda la planilla por fecha (de más antigua a más actual).

Uso:
    python consolidar_facturas.py

Estructura esperada (todo junto a este script):
    FC_30714120421_2501_2512_EL_AMARILLOSRL.xlsx   <- planilla base (original)
    subdiarios/      <- .XLS de los subdiarios mensuales
    comprobantes/    <- .XLS de cada comprobante individual

Genera: planilla_anual.xlsx
"""

import os
import re
import sys
import pandas as pd
from openpyxl import load_workbook
from copy import copy
from glob import glob
from collections import OrderedDict

# ─── CONFIGURACIÓN ────────────────────────────────────────────────────────────
PLANILLA_BASE        = "FC_30714120421_2501_2512_EL_AMARILLOSRL.xlsx"
CARPETA_SUBDIARIOS   = "subdiarios"
CARPETA_COMPROBANTES = "comprobantes"
ARCHIVO_SALIDA       = "planilla_anual.xlsx"
HOJA                 = "FACTURACION"
# ──────────────────────────────────────────────────────────────────────────────

# Mapeo TIPODGI (columna del subdiario) -> código AFIP de 3 dígitos (texto)
MAPEO_TIPODGI = {
    1:  "001",   # Factura A
    2:  "002",   # Nota de Débito A
    3:  "003",   # Nota de Crédito A
    6:  "006",   # Factura B
    7:  "007",   # Nota de Débito B
    8:  "008",   # Nota de Crédito B
    11: "011",   # Factura C
    12: "012",   # Nota de Débito C
    13: "013",   # Nota de Crédito C
    19: "019",   # Factura de Exportación
    20: "020",   # Nota de Débito por op. con el exterior
    21: "021",   # Nota de Crédito por op. con el exterior
}

import unicodedata

# Patrón robusto: identifica el comprobante por el FINAL del nombre
#   (letra + punto de venta - número), sin depender de la palabra "NÚMERO"
#   ni de los acentos. Ejemplos válidos:
#     FACTURA NÚMERO A 0003-00003396.XLS
#     NOTA DE CREDITO NUMERO A 0003-00000226.XLS
PATRON_COLA = re.compile(
    r"(?P<letra>[A-Za-z])[ _](?P<pto>\d+)-(?P<nro>\d+)\.xls$",
    re.IGNORECASE,
)


def _sin_acentos(texto):
    """Quita acentos para comparar de forma robusta."""
    nfkd = unicodedata.normalize("NFD", texto)
    return "".join(c for c in nfkd if unicodedata.category(c) != "Mn")


def _listar_xls(carpeta):
    """Lista archivos .XLS sin duplicados (Windows no distingue may/min)."""
    vistos = {}
    for ruta in glob(os.path.join(carpeta, "*.XLS")) + \
                glob(os.path.join(carpeta, "*.xls")):
        vistos[os.path.normcase(os.path.abspath(ruta))] = ruta
    return sorted(vistos.values())


def codigo_desde_nombre(tipo_word, letra):
    """Convierte el tipo del nombre de archivo + letra al código AFIP."""
    t = tipo_word.upper().replace("_", " ")
    es_credito = "CREDITO" in t
    es_debito  = "DEBITO" in t
    es_factura = ("FACTURA" in t or "EXPORTACION" in t) and not es_credito and not es_debito

    if "EXPORTACION" in t:
        return "019"
    if es_credito:
        return {"A": "003", "B": "008", "C": "013"}.get(letra)
    if es_debito:
        return {"A": "002", "B": "007", "C": "012"}.get(letra)
    if es_factura:
        return {"A": "001", "B": "006", "C": "011", "E": "019"}.get(letra)
    return None


def parsear_nombre(nombre):
    """Devuelve (tipo, pto, nro) a partir del nombre de archivo, o None."""
    m = PATRON_COLA.search(nombre)
    if not m:
        return None
    letra = m.group("letra").upper()
    pto   = int(m.group("pto"))
    nro   = int(m.group("nro"))
    # El tipo está antes de la cola; se compara sin acentos
    cabeza = _sin_acentos(nombre[:m.start()]).upper()
    cod = codigo_desde_nombre(cabeza, letra)
    if cod is None:
        return None
    return (cod, pto, nro)


def indexar_comprobantes(carpeta):
    """Devuelve un dict {(tipo, pto, nro): ruta_xls} con todos los .XLS."""
    indice = {}
    for ruta in _listar_xls(carpeta):
        nombre = os.path.basename(ruta)
        clave = parsear_nombre(nombre)
        if clave is None:
            print(f"  [AVISO] Nombre no reconocido, se omite: {nombre}")
            continue
        indice[clave] = ruta
    return indice


def leer_items_xls(ruta_xls):
    """Lee un .XLS de comprobante y devuelve la lista de ítems."""
    try:
        df = pd.read_excel(ruta_xls, engine="xlrd", header=0)
    except Exception as e:
        print(f"  [ERROR] No se pudo leer {os.path.basename(ruta_xls)}: {e}")
        return None

    if "CANTID" not in df.columns or "DESCRIP" not in df.columns:
        print(f"  [AVISO] {os.path.basename(ruta_xls)} sin columnas esperadas.")
        return None

    items = df[(df["CANTID"] > 0) & (df["DESCRIP"].notna())].copy()
    if items.empty:
        return None

    resultado = []
    for _, item in items.iterrows():
        # Solo el DESCRIP de la fila del ítem (1 renglón), nunca DESCRI1/2/3
        descrip = str(item["DESCRIP"]).strip()

        cod = None
        for campo in ["CODINUEVO", "CODIANTE"]:
            val = item.get(campo)
            if pd.notna(val) and str(val).strip() not in ("", "0", "nan"):
                try:
                    cod = int(float(str(val)))
                    break
                except ValueError:
                    pass

        cant  = int(item["CANTID"])
        neto  = float(item["NETO1"])              # neto real (ya con descuento)
        iva   = float(item["IVA1"])               # IVA real del ítem
        precio_real = neto / cant if cant else 0  # precio unitario con descuento
        total = neto + iva                        # total del concepto

        resultado.append({
            "cod":    cod,
            "desc":   descrip,
            "cant":   cant,
            "precio": precio_real,
            "neto":   neto,
            "iva":    iva,
            "total":  total,
        })
    return resultado


def leer_subdiarios(carpeta):
    """Lee todos los subdiarios y devuelve la lista de comprobantes (cabeceras)."""
    archivos = _listar_xls(carpeta)
    if not archivos:
        return []

    comprobantes = []
    for ruta in archivos:
        print(f"  Leyendo subdiario: {os.path.basename(ruta)}")
        df = pd.read_excel(ruta, engine="xlrd", header=0)
        for _, r in df.iterrows():
            tipodgi = r.get("TIPODGI")
            # Saltear filas sin datos válidos (totales, vacías, etc.)
            if pd.isna(tipodgi):
                continue
            if pd.isna(r.get("COMPROB")) or pd.isna(r.get("SUCURSAL")):
                continue
            cuit_raw = r.get("CUIT")
            if pd.isna(cuit_raw) or str(cuit_raw).strip().lower() in ("", "nan"):
                continue
            try:
                tipo = MAPEO_TIPODGI.get(int(tipodgi))
            except (ValueError, TypeError):
                continue
            if tipo is None:
                print(f"    [AVISO] TIPODGI {tipodgi} no mapeado, fila omitida.")
                continue
            try:
                pto  = int(r["SUCURSAL"])
                nro  = int(r["COMPROB"])
                cuit = int(str(cuit_raw).replace("-", "").strip())
            except (ValueError, TypeError):
                continue

            cai_raw = r.get("CAI")
            try:
                cae = int(cai_raw) if pd.notna(cai_raw) else 0
            except (ValueError, TypeError):
                cae = 0

            comprobantes.append({
                "clave":      (tipo, pto, nro),
                "grupo":      1 if pto in (5, 7) else 3,
                "fecha":      r["FECHA"],
                "tipo":       tipo,
                "pto_venta":  pto,
                "nro_comp":   nro,
                "cuit":       cuit,
                "cliente":    str(r["APEAFIP"]).strip(),
                "total_comp": float(r["TOTAL"]),
                "cae":        cae,
            })
    return comprobantes


def copiar_estilo(origen, destino):
    destino.font          = copy(origen.font)
    destino.alignment     = copy(origen.alignment)
    destino.number_format = origen.number_format


def fila_cabecera(comp):
    """Fila base con los datos de cabecera (detalle vacío)."""
    return {
        1: comp["grupo"], 2: comp["fecha"], 3: comp["tipo"], 4: comp["pto_venta"],
        5: comp["nro_comp"], 6: comp["cuit"], 7: comp["cliente"], 8: comp["total_comp"],
        9: "PESOS", 10: 1, 11: None, 12: "CAE", 13: comp["cae"],
        14: None, 15: None, 16: None, 17: None, 18: None,
        19: None, 20: None, 21: None, 22: None, 23: None, 24: None, 25: None,
    }


def fila_item(comp, item):
    """Fila completa: cabecera + detalle del ítem (todo valores fijos, sin fórmulas)."""
    f = fila_cabecera(comp)
    f.update({
        14: item["cod"], 17: item["desc"], 18: item["cant"], 19: item["precio"],
        20: item["neto"], 21: item["iva"], 25: item["total"],
    })
    return f


def _avisar_faltantes(faltan):
    print(f"\n  ⚠ FALTAN {len(faltan)} comprobante(s) en la carpeta "
          f"'{CARPETA_COMPROBANTES}'.")
    print(f"    Se agregó la fila con la cabecera pero SIN detalle. Faltan:")
    for comp in faltan:
        tipo, pto, nro = comp["clave"]
        fecha = comp["fecha"]
        f_str = fecha.strftime("%d/%m/%Y") if hasattr(fecha, "strftime") else fecha
        print(f"      - tipo {tipo} pto {pto:04d} nro {nro:08d}  ({f_str})  "
              f"{comp['cliente']}")


def main():
    if not os.path.exists(PLANILLA_BASE):
        sys.exit(f"ERROR: No se encontró la planilla base '{PLANILLA_BASE}'.")
    if not os.path.isdir(CARPETA_SUBDIARIOS):
        sys.exit(f"ERROR: No se encontró la carpeta '{CARPETA_SUBDIARIOS}'.")
    if not os.path.isdir(CARPETA_COMPROBANTES):
        sys.exit(f"ERROR: No se encontró la carpeta '{CARPETA_COMPROBANTES}'.")

    print(f"Planilla base: {PLANILLA_BASE}\n")

    print("Leyendo subdiarios...")
    comprobantes = leer_subdiarios(CARPETA_SUBDIARIOS)
    if not comprobantes:
        sys.exit(f"ERROR: No hay subdiarios válidos en '{CARPETA_SUBDIARIOS}'.")
    print(f"  Comprobantes listados: {len(comprobantes)}\n")

    print("Indexando comprobantes (XLS)...")
    indice_xls = indexar_comprobantes(CARPETA_COMPROBANTES)
    print(f"  Archivos XLS reconocidos: {len(indice_xls)}\n")

    wb = load_workbook(PLANILLA_BASE)
    ws = wb[HOJA]

    # Comprobantes ya presentes en la planilla (para no duplicar)
    existentes = set()
    for row in ws.iter_rows(min_row=2):
        if row[4].value is not None:
            existentes.add((str(row[2].value), row[3].value, row[4].value))

    # Última fila con datos (referencia de estilo)
    fila_ref = 1
    for row in ws.iter_rows(min_row=2):
        if row[4].value is not None:
            fila_ref = row[0].row

    # Procesar cada comprobante del subdiario
    bloques_nuevos = []   # cada uno: {fecha, filas:[dict col->val]}
    faltan = []
    ya_cargados = 0
    n_items = 0

    for comp in comprobantes:
        if comp["clave"] in existentes:
            ya_cargados += 1
            continue
        ruta = indice_xls.get(comp["clave"])
        items = leer_items_xls(ruta) if ruta else None

        if items:
            filas = [fila_item(comp, it) for it in items]
            n_items += len(filas)
        else:
            # Falta el XLS (o no tiene ítems): fila solo con cabecera
            filas = [fila_cabecera(comp)]
            faltan.append(comp)

        bloques_nuevos.append({"fecha": comp["fecha"], "filas": filas})

    nuevos = len(comprobantes) - ya_cargados
    print(f"Comprobantes ya cargados (omitidos): {ya_cargados}")
    print(f"Comprobantes nuevos:                 {nuevos}")
    print(f"  - completos (con XLS):             {nuevos - len(faltan)}")
    print(f"  - sin XLS (cabecera vacía):        {len(faltan)}")
    print(f"Filas de ítems agregadas:            {n_items}")

    if not bloques_nuevos:
        print("\nNo hay comprobantes nuevos para agregar.")
        return

    # ── Reconstruir la hoja: bloques existentes + nuevos, ordenados por fecha ──
    bloques = []
    orden = 0
    clave_ant = None
    bloque = None
    for row in ws.iter_rows(min_row=2):
        if row[4].value is None:
            continue
        clave = (str(row[2].value), row[3].value, row[4].value)
        vals = {c: ws.cell(row=row[0].row, column=c).value for c in range(1, 26)}
        if clave != clave_ant:
            bloque = {"fecha": row[1].value, "orden": orden, "filas": [], "nuevo": False}
            bloques.append(bloque)
            orden += 1
            clave_ant = clave
        bloque["filas"].append(vals)

    for bn in bloques_nuevos:
        bloques.append({"fecha": bn["fecha"], "orden": orden,
                        "filas": bn["filas"], "nuevo": True})
        orden += 1

    bloques.sort(key=lambda b: (b["fecha"], b["orden"]))

    # ── Reescribir la hoja ────────────────────────────────────────────────────
    fila = 2
    for b in bloques:
        for vals in b["filas"]:
            for col in range(1, 26):
                v = vals[col]
                cell = ws.cell(row=fila, column=col)
                cell.value = v
                if b["nuevo"] and fila_ref:
                    copiar_estilo(ws.cell(row=fila_ref, column=col), cell)
                    if col == 2:
                        cell.number_format = "dd/mm/yyyy"
            fila += 1

    wb.save(ARCHIVO_SALIDA)
    print(f"\n✓ Listo. Archivo guardado como: {ARCHIVO_SALIDA}")

    ncnd = [b for b in bloques_nuevos
            if b["filas"] and str(b["filas"][0][3]) in ("002", "003")]
    if ncnd:
        print(f"\n  ATENCIÓN: agregaste {len(ncnd)} nota(s) de crédito/débito.")
        print(f"  Recordá completar a mano la columna 'Imputación' (col K).")

    if faltan:
        _avisar_faltantes(faltan)


if __name__ == "__main__":
    main()
